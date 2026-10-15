from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / '第六节课AI硬件作业-飞书文档正文.md'
OUTPUT = ROOT / 'docs' / '第六节课AI硬件作业-含截图导入版.docx'


def set_font(style, name='Microsoft YaHei', size=10.5, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{attr}'), name)


def format_inline(paragraph, text):
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
        elif part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.9)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

styles = doc.styles
set_font(styles['Normal'])
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.22
set_font(styles['Title'], size=19, bold=True)
styles['Title'].paragraph_format.space_after = Pt(16)
set_font(styles['Heading 1'], size=14, bold=True)
styles['Heading 1'].paragraph_format.space_before = Pt(16)
styles['Heading 1'].paragraph_format.space_after = Pt(8)
set_font(styles['Heading 2'], size=11.5, bold=True)
styles['Heading 2'].paragraph_format.space_before = Pt(11)
styles['Heading 2'].paragraph_format.space_after = Pt(5)
set_font(styles['List Bullet'])
set_font(styles['List Number'])

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    raw = lines[i].strip()
    if not raw:
        i += 1
        continue
    if raw.startswith('```'):
        code = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith('```'):
            code.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run('\n'.join(code))
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        i += 1
        continue
    if raw.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            parts = [v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', v or '') for v in parts):
                rows.append(parts)
            i += 1
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci, value in enumerate(row):
                    cell = table.cell(ri, ci)
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if ri == 0:
                        shade(cell, 'E8EEF5')
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        continue
    image_match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', raw)
    if image_match:
        caption, rel = image_match.groups()
        path = SOURCE.parent / rel
        if path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            with Image.open(path) as im:
                w, h = im.size
            max_w, max_h = Cm(16.2), Cm(17.0)
            ratio = min(max_w / w, max_h / h)
            p.add_run().add_picture(str(path), width=int(w * ratio), height=int(h * ratio))
            cp = doc.add_paragraph(caption)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(12)
            for run in cp.runs:
                run.font.size = Pt(9)
        i += 1
        continue
    if raw.startswith('# '):
        doc.add_paragraph(raw[2:].replace('｜', ' '), 'Title')
    elif raw.startswith('## '):
        doc.add_paragraph(raw[3:].replace('｜', ' '), 'Heading 1')
    elif raw.startswith('### '):
        doc.add_paragraph(raw[4:].replace('｜', ' '), 'Heading 2')
    elif raw.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        format_inline(p, raw[2:])
    elif re.match(r'^\d+\. ', raw):
        p = doc.add_paragraph(style='List Number')
        format_inline(p, re.sub(r'^\d+\. ', '', raw))
    else:
        p = doc.add_paragraph()
        format_inline(p, raw)
    i += 1

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
